# -*- coding: utf-8 -*-
import docx, sys
from docx.shared import Pt, RGBColor
sys.stdout.reconfigure(encoding='utf-8')

PATH = 'D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript/RCDome_manuscript_V024_supplementary_information.docx'
doc = docx.Document(PATH)

# ---- 1. Style definitions (uniform Arial) ----
def set_style(name, size, bold):
    try:
        st = doc.styles[name]
        st.font.name = 'Arial'
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    except KeyError:
        pass

set_style('Normal', 11, False)
set_style('Title', 16, True)
set_style('Heading 1', 14, True)
set_style('Heading 2', 12, True)

# ---- 2. Uniformize paragraph runs ----
for p in doc.paragraphs:
    style_name = p.style.name
    target_size = {'Title': 16, 'Heading 1': 14, 'Heading 2': 12}.get(style_name, 11)
    is_note_title = p.text.strip().startswith('RCDome:')
    for r in p.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(target_size)
        r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        if is_note_title:
            r.font.bold = True
            r.font.size = Pt(13)

# ---- 3. Uniformize table runs ----
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    # header row bold
    try:
        for cell in t.rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
    except Exception:
        pass

doc.save(PATH)
print('uniformized: Arial, body 11pt, headings 14/12, title 16, tables 10pt, note titles bold 13')

# -*- coding: utf-8 -*-
import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

path = 'manuscript_V027/RCDome_manuscript_V027_supplementary_information.docx'
d = docx.Document(path)

def set_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)

rep81 = [
    ("Four evidence tiers govern", "Four provenance categories govern"),
    ("(4) the output ends with the exact verbatim disclaimer",
     "(4) the exact verbatim disclaimer is appended at the very end"),
    ('(5) the output ends with exactly three numbered "Suggested Queries:"',
     '(5) the synthesis body ends with exactly three numbered "Suggested Queries:" immediately before the disclaimer'),
]
rep93 = [
    ("The 100% positive and negative results, with zero hallucinated PMIDs, demonstrate that full coverage and zero hallucination are compatible in the same module.",
     "The 100% positive and negative results recorded zero hallucinated PMIDs across the benchmark, demonstrating that full coverage and zero observed hallucination were achieved together in the same module."),
]

done1 = done3 = False
for p in d.paragraphs:
    if 'Four evidence tiers govern' in p.text:
        t = p.text
        for a,b in rep81:
            t = t.replace(a,b)
        set_text(p, t)
        done1 = True
    elif 'zero hallucination are compatible' in p.text:
        t = p.text
        for a,b in rep93:
            t = t.replace(a,b)
        set_text(p, t)
        done3 = True

assert done1 and done3, (done1, done3)
d.save(path)
print('docx updated: Higor points 1/2/4 applied (provenance categories, ends-with order, zero observed hallucination)')

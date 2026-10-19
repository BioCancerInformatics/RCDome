# -*- coding: utf-8 -*-
"""
Fix script — Victor "RELATÓRIO DE VERIFICAÇÃO DO APP" (approved 2026-09-19).
1. "back to the dome" -> click parent Home tab
2. mod_od.R -> composite founding_pmid renders real PMIDs
3. decoder -> drop 29362479 from primary_refs_pmids for 8 NOT_IN_NCCD_2018 entities
4. decoder -> VCD citation clarified
5. mod_rcd_grid.R -> collapse = "\n\n"
"""
import io, csv, os, sys, shutil, datetime
sys.stdout.reconfigure(encoding='utf-8')

ROOT = "D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript"
GH   = "D:/RCDome_github"
log  = []

def note(msg):
    log.append(msg)
    print(msg)

HOME_CLICK = ('onclick="(function(){var l=window.parent.document.querySelectorAll(\'a[data-toggle=tab]\');'
              'for(var i=0;i<l.length;i++){if(l[i].textContent.replace(/\\s+/g,\' \').trim().indexOf(\'Home\')===0)'
              '{l[i].click();return;}}})();return false;"')

# ───────────────────────── Fix 1: HTML back-links ─────────────────────────
html_files = [
    os.path.join(ROOT, "RCDome_distribution/www/explore.html"),
    os.path.join(ROOT, "RCDome_distribution/www/detail_page_wireframe.html"),
    os.path.join(ROOT, "RCDome_distribution/www/module_placeholder.html"),
    os.path.join(GH, "app/www/explore.html"),
    os.path.join(GH, "app/www/detail_page_wireframe.html"),
    os.path.join(GH, "app/www/module_placeholder.html"),
    os.path.join(GH, "app/explore.html"),
    os.path.join(GH, "app/detail_page_wireframe.html"),
]
for f in html_files:
    if not os.path.exists(f):
        note(f"  [SKIP] {f} not found"); continue
    s = io.open(f, encoding='utf-8').read()
    n = s.count('href="landing_page.html"')
    s2 = s.replace('href="landing_page.html"', 'href="#" ' + HOME_CLICK)
    io.open(f, 'w', encoding='utf-8').write(s2)
    note(f"  [HTML] {f}: {n} back-link(s) -> Home-tab click")

# generate_detail_page.py (the template that produces detail_page_wireframe.html)
gp = os.path.join(GH, "app/generate_detail_page.py")
if os.path.exists(gp):
    s = io.open(gp, encoding='utf-8').read()
    n = s.count('href="landing_page.html"')
    s2 = s.replace('href="landing_page.html"', 'href="#" ' + HOME_CLICK)
    io.open(gp, 'w', encoding='utf-8').write(s2)
    note(f"  [PY]  {gp}: {n} back-link(s) -> Home-tab click")

# ───────────────────────── Fix 2: mod_od.R ─────────────────────────
mod_od_old = '''          } else if (f %in% c("founding_pmid", "reference_pmid", "primary_discovery_reference")) {
            val <- paste0("<a href='https://pubmed.ncbi.nlm.nih.gov/", as.character(v),
                          "/' target='_blank'>PMID ", as.character(v), "</a>")
          } else if (f == "reference_citation") {
            pmid <- r$reference_pmid %||% ""
            if (!nzchar(pmid)) pmid <- r$founding_pmid %||% ""
            if (nzchar(pmid)) {
              val <- paste0("<a href='https://pubmed.ncbi.nlm.nih.gov/", pmid,
                            "/' target='_blank'>", as.character(v), "</a>")
            } else {
              val <- as.character(v)
            }
          }'''

mod_od_new = '''          } else if (f == "founding_pmid") {
            fv <- as.character(v)
            if (grepl("[0-9]", fv) && !grepl("[^0-9; ]", fv)) {
              val <- render_pmids(fv)
              if (!nzchar(val)) return(NULL)
            } else {
              rp <- as.character(r$reference_pmid %||% "")
              links <- render_pmids(rp)
              val <- if (nzchar(links)) paste0("Composite — ", links) else fv
            }
          } else if (f %in% c("reference_pmid", "primary_discovery_reference")) {
            val <- render_pmids(v)
            if (!nzchar(val)) return(NULL)
          } else if (f == "reference_citation") {
            pmid <- as.character(r$reference_pmid %||% "")
            if (!nzchar(pmid)) pmid <- as.character(r$founding_pmid %||% "")
            if (nzchar(pmid) && grepl("[0-9]", pmid)) {
              first <- trimws(strsplit(pmid, ";")[[1]])[1]
              val <- paste0("<a href='https://pubmed.ncbi.nlm.nih.gov/", first,
                            "/' target='_blank'>", as.character(v), "</a>")
            } else {
              val <- as.character(v)
            }
          }'''

for f in [os.path.join(ROOT, "RCDome_distribution/modules/mod_od.R"),
          os.path.join(GH, "app/modules/mod_od.R")]:
    s = io.open(f, encoding='utf-8').read()
    if mod_od_old not in s:
        note(f"  [WARN] mod_od.R block not found in {f}")
        continue
    s2 = s.replace(mod_od_old, mod_od_new)
    io.open(f, 'w', encoding='utf-8').write(s2)
    note(f"  [R]    {f}: composite PMID render fixed")

# ───────────────────────── Fix 5: mod_rcd_grid.R ─────────────────────────
for f in [os.path.join(ROOT, "RCDome_distribution/modules/mod_rcd_grid.R"),
          os.path.join(GH, "app/modules/mod_rcd_grid.R")]:
    s = io.open(f, encoding='utf-8').read()
    old = '      ), collapse = "\\n")\n'
    # only the build_decoder_desc join (line ~59) — it is the FIRST occurrence
    idx = s.find(old)
    if idx < 0:
        note(f"  [WARN] mod_rcd_grid.R collapse line not found in {f}")
        continue
    s2 = s[:idx] + '      ), collapse = "\\n\\n")\n' + s[idx+len(old):]
    io.open(f, 'w', encoding='utf-8').write(s2)
    note(f"  [R]    {f}: build_decoder_desc collapse -> \\n\\n")

# ───────────────────────── Fix 3 + 4: decoder data ─────────────────────────
DROP_29362479 = {
    "Apoptosis", "Autophagy", "Methuosis", "Mitochondrial permeability transition",
    "Mitoptosis", "Necrosis", "Oxeiptosis", "Paraptosis",
}
VCD_CITATION = "Kosta et al. 2004 \u2014 name source; describes a non-vacuolar pathway"

src_csv = os.path.join(ROOT, "RCDome Decoder/RCDome_Decoder_53.csv")
src_xlsx = os.path.join(ROOT, "RCDome Decoder/RCDome_Decoder_53.xlsx")

with io.open(src_csv, encoding='utf-8-sig', newline='') as fh:
    reader = csv.DictReader(fh)
    fieldnames = reader.fieldnames
    rows = list(reader)

changed = 0
for r in rows:
    nm = r.get("rcd_form", "")
    if nm in DROP_29362479:
        refs = r.get("primary_refs_pmids", "")
        parts = [p.strip() for p in refs.split(";") if p.strip() and p.strip() != "29362479"]
        new = ";".join(parts)
        if new != refs:
            r["primary_refs_pmids"] = new
            note(f"  [DATA] {nm}: primary_refs_pmids {refs!r} -> {new!r}")
            changed += 1
    if nm == "Vacuolar cell death":
        old = r.get("reference_citation", "")
        r["reference_citation"] = VCD_CITATION
        note(f"  [DATA] Vacuolar cell death: reference_citation {old!r} -> {VCD_CITATION!r}")
        changed += 1

if changed:
    with io.open(src_csv, 'w', encoding='utf-8-sig', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    note(f"  [DATA] wrote {src_csv} ({changed} cell changes)")

    # XLSX — edit cells in place (preserves formatting)
    import openpyxl
    wb = openpyxl.load_workbook(src_xlsx)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    col_form = hdr.index("rcd_form")
    col_prim = hdr.index("primary_refs_pmids")
    col_cit  = hdr.index("reference_citation")
    for row in ws.iter_rows(min_row=2):
        nm = row[col_form].value
        if nm in DROP_29362479:
            cell = row[col_prim]
            parts = [p.strip() for p in str(cell.value).split(";") if p.strip() and p.strip() != "29362479"]
            cell.value = ";".join(parts)
        if nm == "Vacuolar cell death":
            row[col_cit].value = VCD_CITATION
    wb.save(src_xlsx)
    note(f"  [DATA] wrote {src_xlsx}")
else:
    note("  [DATA] no decoder cell changes detected — check!")

# propagate decoder CSV/XLSX to app + GitHub copies
copies = [
    (src_csv, os.path.join(ROOT, "RCDome_distribution/data/RCDome_Decoder_53.csv")),
    (src_csv, os.path.join(GH, "data/RCDome_Decoder_53.csv")),
    (src_csv, os.path.join(GH, "app/data/RCDome_Decoder_53.csv")),
    (src_xlsx, os.path.join(GH, "data/RCDome_Decoder_53.xlsx")),
]
for s, d in copies:
    shutil.copy2(s, d)
    note(f"  [COPY] {s} -> {d}")

# ───────────────────────── log file ─────────────────────────
log_path = os.path.join(ROOT, "docs/TRACEABILITY_Victor_app_fixes_2026-09-19.md")
with io.open(log_path, 'w', encoding='utf-8') as fh:
    fh.write("# Traceability — Victor app-report fixes (2026-09-19)\n\n")
    fh.write("Approved fixes applied. No data deleted.\n\n")
    fh.write("\n".join("- " + m for m in log))
note(f"  [LOG]  {log_path}")
print("\nDONE —", len(log), "log entries")

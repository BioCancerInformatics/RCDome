# -*- coding: utf-8 -*-
"""FINAL deep compliance audit — Victor app-report vs app code, supplementary tables, matrix, decoder."""
import csv, io, os, sys, hashlib, subprocess
sys.stdout.reconfigure(encoding='utf-8')
csv.field_size_limit(10**9)

ROOT = "D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript"
GH   = "D:/RCDome_github"
res = []
def ok(m): res.append(("PASS", m)); print(f"  [PASS] {m}")
def bad(m): res.append(("FAIL", m)); print(f"  [FAIL] {m}")
def md5(p):
    try: return hashlib.md5(open(p,'rb').read()).hexdigest()
    except Exception: return None
def read(p):
    return io.open(p, encoding='utf-8').read()

# ───────────────────────── A. DECODER ─────────────────────────
print("="*78); print("A. DECODER (CSV + XLSX)"); print("="*78)
src_csv = os.path.join(ROOT,"RCDome Decoder/RCDome_Decoder_53.csv")
rows = list(csv.DictReader(io.open(src_csv, encoding='utf-8-sig')))
ok(f"53 rows") if len(rows)==53 else bad(f"rows={len(rows)}")
ok(f"39 cols") if len(rows[0])==39 else bad(f"cols={len(rows[0])}")
by = {r['rcd_form']: r for r in rows}
DROP = {"Apoptosis","Autophagy","Methuosis","Mitochondrial permeability transition","Mitoptosis","Necrosis","Oxeiptosis","Paraptosis"}
for e in DROP:
    v = by[e].get('primary_refs_pmids','')
    ok(f"{e}: primary_refs={v}") if "29362479" not in v else bad(f"{e} -> {v}")
c = by.get("Senescence and Autophagy in Cancer", {})
for k,exp in [("category","COMPOSITE"),("definition_status","COMPOSITE"),("galluzzi_nccd_status","COMPOSITE")]:
    ok(f"composite {k}={exp}") if c.get(k)==exp else bad(f"composite {k}={c.get(k)!r}")
ok("composite citation explains NCCD component") if ('NCCD_2018' in c.get('reference_citation','') and 'not a distinct NCCD entity' in c.get('reference_citation','')) else bad("composite citation")
v = by.get("Vacuolar cell death", {})
ok("VCD category=TAXON") if v.get('category')=='TAXON' else bad(f"VCD category={v.get('category')}")
ok("VCD founding=15358773") if v.get('founding_pmid')=='15358773' else bad("VCD founding")
ok("VCD secondary=24344187;35390283") if v.get('secondary_refs_pmids')=='24344187;35390283' else bad("VCD secondary")
ok("VCD citation = Dictyostelium + plant") if ('Dictyostelium' in v.get('reference_citation','') and '24344187' in v.get('reference_citation','') and '35390283' in v.get('reference_citation','')) else bad("VCD citation")
comp_others = [r['rcd_form'] for r in rows if r['galluzzi_nccd_status']=='COMPOSITE']
ok("COMPOSITE status only on composite") if comp_others==["Senescence and Autophagy in Cancer"] else bad(f"COMPOSITE: {comp_others}")
bad_ones = [r['rcd_form'] for r in rows if r['galluzzi_nccd_status']=='NOT_IN_NCCD_2018' and r['rcd_form']!='Senescence and Autophagy in Cancer' and '29362479' in r['primary_refs_pmids']]
ok("no NOT_IN_NCCD (non-composite) still lists 29362479") if not bad_ones else bad(f"{bad_ones}")

# decoder XLSX cells
import openpyxl
wb = openpyxl.load_workbook(os.path.join(ROOT,"RCDome Decoder/RCDome_Decoder_53.xlsx"))
ws = wb.active; hdr=[c.value for c in ws[1]]
ix = {h:i for i,h in enumerate(hdr)}
xlsx_ok = True
for row in ws.iter_rows(min_row=2, values_only=True):
    nm = row[ix['rcd_form']]
    if nm in DROP and '29362479' in str(row[ix['primary_refs_pmids']]): xlsx_ok=False; bad(f"XLSX {nm} still 29362479")
    if nm=="Senescence and Autophagy in Cancer" and row[ix['galluzzi_nccd_status']]!="COMPOSITE": xlsx_ok=False; bad("XLSX composite status")
    if nm=="Vacuolar cell death" and 'Dictyostelium' not in str(row[ix['reference_citation']]): xlsx_ok=False; bad("XLSX VCD citation")
if xlsx_ok: ok("decoder XLSX cells consistent")

# ───────────────────────── B. APP CODE ─────────────────────────
print("="*78); print("B. APP CODE (R modules)"); print("="*78)
mo = read(os.path.join(ROOT,"RCDome_distribution/modules/mod_od.R"))
ok("mod_od.R composite founding branch") if ('f == "founding_pmid"' in mo and 'Composite —' in mo) else bad("mod_od composite")
ok("mod_od.R multi-PMID render_pmids") if ('val <- render_pmids(v)' in mo) else bad("mod_od render_pmids")
ok("mod_od.R citation first PMID") if ('first <- trimws(strsplit(pmid, ";")[[1]])[1]' in mo) else bad("mod_od first")
mg = read(os.path.join(ROOT,"RCDome_distribution/modules/mod_rcd_grid.R"))
ok('mod_rcd_grid.R collapse="\\n\\n"') if 'collapse = "\\n\\n"' in mg else bad("mod_rcd_grid collapse")
mb = read(os.path.join(ROOT,"RCDome_distribution/modules/mod_board.R"))
ok('mod_board.R NCCD desc includes / COMPOSITE') if 'NOT_IN_NCCD_2018 / COMPOSITE' in mb else bad("mod_board desc")

R = r'"C:\Program Files\R\R-4.2.3\bin\Rscript.exe"'
for m in ["mod_od.R","mod_rcd_grid.R","mod_board.R"]:
    p = os.path.join(ROOT,"RCDome_distribution/modules",m).replace("\\","/")
    r = subprocess.run(f'{R} -e "invisible(parse(\'{p}\'))"', shell=True, capture_output=True, text=True)
    ok(f"{m} parses") if "rror" not in r.stderr+r.stdout else bad(f"{m}: {r.stderr[:150]}")

# ───────────────────────── C. HTML NAVIGATION ─────────────────────────
print("="*78); print("C. HTML 'Back to the dome' (unified, all copies)"); print("="*78)
uni = "getElementById('home_iframe')"
files = {
 "detail_page_wireframe.html": [os.path.join(ROOT,"RCDome_distribution/detail_page_wireframe.html"), os.path.join(ROOT,"RCDome_distribution/www/detail_page_wireframe.html"), os.path.join(GH,"app/detail_page_wireframe.html"), os.path.join(GH,"app/www/detail_page_wireframe.html")],
 "explore.html": [os.path.join(ROOT,"RCDome_distribution/explore.html"), os.path.join(ROOT,"RCDome_distribution/www/explore.html"), os.path.join(GH,"app/explore.html"), os.path.join(GH,"app/www/explore.html")],
 "module_placeholder.html": [os.path.join(ROOT,"RCDome_distribution/www/module_placeholder.html"), os.path.join(GH,"app/www/module_placeholder.html")],
 "generate_detail_page.py": [os.path.join(ROOT,"RCDome_distribution/generate_detail_page.py"), os.path.join(GH,"app/generate_detail_page.py")],
}
allc = True
for name, paths in files.items():
    vals=[md5(p) for p in paths]
    same = len(set(vals))==1
    s = read(paths[0])
    has_uni = uni in s; no_href = 'href="landing_page.html"' not in s
    ok(f"{name}: copies-identical={same}, unified={has_uni}, no-stale-href={no_href}") if (same and has_uni and no_href) else bad(f"{name}: same={same} uni={has_uni} href={no_href}")
    if not (same and has_uni and no_href): allc=False

# ───────────────────────── D. SUPPLEMENTARY TABLES ─────────────────────────
print("="*78); print("D. SUPPLEMENTARY TABLES (Excel)"); print("="*78)
wb = openpyxl.load_workbook(os.path.join(ROOT,"manuscript_V026/RCDome_manuscript_V026_supplementary_tables.xlsx"))
ws3 = wb["Table S3"]; good3=True
for r in ws3.iter_rows(min_row=3, values_only=True):
    if r[1] in DROP and '29362479' in str(r[6]): good3=False; bad(f"Table S3 {r[1]} still 29362479")
ok("Table S3: 8 entities Key References cleaned") if good3 else None
ws11 = wb["Table S11"]; h=[c.value for c in ws11[1]]; jx={x:i for i,x in enumerate(h)}
good11=True
for r in ws11.iter_rows(min_row=2, values_only=True):
    nm=r[jx['rcd_form']]
    if nm in DROP and '29362479' in str(r[jx['primary_refs_pmids']]): good11=False; bad(f"S11 {nm} primary_refs")
    if nm=="Senescence and Autophagy in Cancer":
        if r[jx['galluzzi_nccd_status']]!="COMPOSITE": good11=False; bad("S11 composite status")
        if 'not a distinct NCCD entity' not in str(r[jx['reference_citation']]): good11=False; bad("S11 composite citation")
    if nm=="Vacuolar cell death" and 'Dictyostelium' not in str(r[jx['reference_citation']]): good11=False; bad("S11 VCD citation")
ok("Table S11: all decoder changes mirrored") if good11 else None

# ───────────────────────── E. MATRIX ─────────────────────────
print("="*78); print("E. MATRIX (V026)"); print("="*78)
mp = os.path.join(ROOT,"output databases/Phase 2C output/unified_gene_rcd_matrix_V026.csv")
with io.open(mp, encoding='utf-8-sig') as f:
    mh = next(csv.reader(f))
mirrors = [c for c in mh if c in ("primary_refs_pmids","reference_citation","galluzzi_nccd_status","founding_pmid")]
ok("matrix does NOT mirror decoder reference fields") if not mirrors else bad(f"matrix mirrors: {mirrors}")
if 'decoder_category' in mh:
    mrows = list(csv.DictReader(io.open(mp, encoding='utf-8-sig')))
    vcd_cat = set(r['decoder_category'] for r in mrows if r.get('rcd_form_canonical')=='Vacuolar cell death')
    ok(f"matrix VCD decoder_category=TAXON") if vcd_cat=={'TAXON'} else bad(f"matrix VCD={vcd_cat}")
else:
    bad("matrix missing decoder_category")

# ───────────────────────── F. MANUSCRIPT ─────────────────────────
print("="*78); print("F. MANUSCRIPT (docx)"); print("="*78)
import docx
d = docx.Document(os.path.join(ROOT,"manuscript_V026/RCDome_manuscript_governance_V026.docx"))
paras = [p.text for p in d.paragraphs]
ok("no 'fixedfixed'") if not any('fixedfixed' in t for t in paras) else bad("still fixedfixed")
ok("'is fixed and' present (typo fixed)") if any('is fixed and is not modified' in t for t in paras) else bad("typo fix missing")
# the only 29362479 occurrences must be the NCCD consensus citation
legit = all(('Galluzzi' in t and 'NCCD' in t) or 'PMID 29362479' in t for t in paras if '29362479' in t)
ok("manuscript 29362479 only as NCCD consensus citation") if legit else bad("manuscript 29362479 usage")

# ───────────────────────── G. PROPAGATION (md5) ─────────────────────────
print("="*78); print("G. PROPAGATION (md5 across copies)"); print("="*78)
pairs = [
 ("decoder CSV", os.path.join(ROOT,"RCDome Decoder/RCDome_Decoder_53.csv"), [os.path.join(ROOT,"RCDome_distribution/data/RCDome_Decoder_53.csv"), os.path.join(GH,"data/RCDome_Decoder_53.csv"), os.path.join(GH,"app/data/RCDome_Decoder_53.csv")]),
 ("decoder XLSX", os.path.join(ROOT,"RCDome Decoder/RCDome_Decoder_53.xlsx"), [os.path.join(GH,"data/RCDome_Decoder_53.xlsx")]),
 ("mod_od.R", os.path.join(ROOT,"RCDome_distribution/modules/mod_od.R"), [os.path.join(GH,"app/modules/mod_od.R")]),
 ("mod_rcd_grid.R", os.path.join(ROOT,"RCDome_distribution/modules/mod_rcd_grid.R"), [os.path.join(GH,"app/modules/mod_rcd_grid.R")]),
 ("mod_board.R", os.path.join(ROOT,"RCDome_distribution/modules/mod_board.R"), [os.path.join(GH,"app/modules/mod_board.R")]),
]
for name, s, ds in pairs:
    m0=md5(s); allok=all(md5(d)==m0 for d in ds)
    ok(f"{name}: {len(ds)} copies identical") if allok else bad(f"{name} mismatch")

# ───────────────────────── H. BOARD PROMPT ─────────────────────────
print("="*78); print("H. BOARD PROMPT"); print("="*78)
bp = read(os.path.join(ROOT,"RCDome_distribution/board/RCD_Board_DeepSeek_Prompt.md"))
ok("no hardcoded stale PMIDs") if ('29362479' not in bp and '15358773' not in bp) else bad("prompt hardcodes PMIDs")

nfail = sum(1 for s,_ in res if s=="FAIL")
print("="*78); print(f"RESULT: {len(res)} checks, {nfail} FAIL"); print("="*78)
with io.open(os.path.join(ROOT,"docs/AUDIT_Victor_compliance_FINAL_2026-09-19.md"),'w',encoding='utf-8') as f:
    f.write("# FINAL compliance audit — Victor app-report (2026-09-19)\n\n")
    f.write(f"**Result: {len(res)} checks, {nfail} FAIL**\n\n")
    f.write("Scope: app code, supplementary tables, matrix, decoder, manuscript, board prompt, propagation.\n\n")
    for s,m in res: f.write(f"- [{s}] {m}\n")
print("saved docs/AUDIT_Victor_compliance_FINAL_2026-09-19.md")

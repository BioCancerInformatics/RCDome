# -*- coding: utf-8 -*-
"""Fetch verbatim sentences for the agreement pool (Option B verbatim-grounded upgrade).
Reads the matrix, builds the fetch queue (agreement pairs lacking elicit/deep verbatim),
fetches abstracts via Europe PMC by DOI, extracts the sentence linking gene->form,
and saves progress incrementally to JSON. Resumable.
"""
import sys, csv, glob, re, json, time, os
import requests
import openpyxl
csv.field_size_limit(10**9)

sys.stdout.reconfigure(encoding='utf-8')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
base = glob.glob(os.path.join(ROOT, 'manuscript_V027', 'Higor*'))[0] + os.sep
MATRIX = os.path.join(ROOT, 'RCDome_distribution', 'data', 'unified_gene_rcd_matrix_V027.csv')
PROGRESS = os.path.join(ROOT, 'manuscript_V027', 'verbatim_agreement_progress.json')
OUT = os.path.join(ROOT, 'manuscript_V027', 'verbatim_agreement_fetched.csv')

ALIAS = {
 'RELA': r'nf-?[kκ]b|nf-?kappa\s*b|p65', 'LGALS1': 'galectin-?1', 'PLAGL1': 'zac1', 'MAPK8': 'jnk1?',
 'PPARGC1A': r'pgc-?1[αa]', 'BECN1': 'beclin[ -]?1', 'SQSTM1': 'p62', 'NFE2L2': 'nrf2', 'TP53': 'p53',
 'CDKN1A': 'p21', 'CDKN2A': 'p16', 'GSDMD': 'gasdermin[ -]?d', 'RIPK1': 'rip1', 'RIPK3': 'rip3',
 'ELAVL1': 'hur', 'SERPINF1': 'pedf', 'PDCD6IP': 'alix', 'STK11': 'lkb1', 'MAP3K5': 'ask1', 'SFN': '14-3-3|stratifin',
 'DAP': 'dapk|dap-kinase', 'GHITM': 'tmbim', 'NOL3': 'arc', 'TEK': 'tie2', 'FAS': 'cd95', 'BAX': 'bax', 'BAK1': 'bak',
 'PRKAA1': 'ampk', 'PRKAA2': 'ampk', 'PPARG': r'ppar[γg]', 'HIF1A': r'hif-?1[αa]', 'AKT1': 'akt', 'MTOR': 'mtor',
 'MAPK1': 'erk2?', 'MAPK3': 'erk1', 'MAPK14': r'p38[αa]?', 'CASP3': 'caspase[ -]?3', 'CASP1': 'caspase[ -]?1',
 'TNF': r'tnf-?[αa]', 'IL1B': r'il-?1[βb]', 'IL18': 'il-?18', 'FOXM1': 'foxm1', 'FOXO3': 'foxo3a?',
 'SNCA': r'[αa]-synuclein', 'APP': 'amyloid precursor', 'TARDBP': 'tdp-?43', 'P4HB': 'pdi|protein disulfide isomerase',
 'MAP3K7': 'tak1', 'NFKB1': r'nf-?[kκ]b', 'NFKBIA': r'nf-?[kκ]b',
}
FORM_TOKENS = {
 'Apoptosis': 'apoptos|fas-mediated', 'Ferroptosis': 'ferroptos', 'Autophagy': 'autophag',
 'Pyroptosis': 'pyroptos', 'Necrosis': 'necrosis|necrotic', 'Necroptosis': 'necroptos',
 'Anoikis': 'anoikis', 'Cellular senescence': 'senescen', 'Autosis': 'autosis', 'Cuproptosis': 'cuproptos',
 'Mitotic catastrophe': 'mitotic', 'Mitotic death': 'mitotic', 'Parthanatos': 'parthanatos',
 'MPT-driven necrosis': 'permeability transition|mptp|mpt', 'Immunogenic cell death': 'immunogenic',
 'Programmed cell death': 'programmed cell death', 'Regulated Necrosis': 'regulated necrosis',
 'Senescence and Autophagy in Cancer': 'senescen|autophag', 'Disulfidptosis': 'disulfidptosis',
 'Efferocytosis': 'efferocytosis', 'PANoptosis': 'panoptos', 'Entotic cell death': 'entosis', 'NETotic cell death': 'netosis',
}
def gene_pat(g):
    return ALIAS.get(g.upper(), r'\b' + re.escape(g.lower()) + r'\b')
def form_pat(f):
    return FORM_TOKENS.get(f, re.escape(f.split('(')[0].strip().lower()))

def clean_html(s):
    s = re.sub(r'<[^>]+>', ' ', s or '')
    s = re.sub(r'\s+', ' ', s)
    return s.strip()

def fetch_abstract(doi, sess):
    url = f'https://www.ebi.ac.uk/europepmc/webservices/rest/search?query=DOI:"{doi}"&resultType=core&format=json'
    for attempt in range(3):
        try:
            r = sess.get(url, timeout=25)
            if r.status_code == 200:
                d = r.json()
                if d.get('hitCount', 0) > 0:
                    return d['resultList']['result'][0].get('abstractText', '') or ''
                return ''
            time.sleep(1.5 + attempt)
        except Exception:
            time.sleep(1.5 + attempt)
    return '__ERR__'

def extract_verbatim(abstract, gene, form):
    text = clean_html(abstract)
    if not text:
        return ''
    gpat = re.compile(gene_pat(gene), re.I)
    fpat = re.compile(form_pat(form), re.I)
    sentences = re.split(r'(?<=[.!?])\s+', text)
    hits = [s.strip() for s in sentences if gpat.search(s) and fpat.search(s)]
    if hits:
        return hits[0][:600]
    return ''

def main(limit=None):
    r = list(csv.DictReader(open(MATRIX, encoding='utf-8-sig')))
    # track pairs (already verbatim-grounded)
    track = set()
    for t in ['A', 'B']:
        wb = openpyxl.load_workbook(base + f'human_resolution_worksheet_track{t}_Higor_Victor.xlsx', read_only=True)
        ws = wb.active; hdr = [c.value for c in ws[1]]; cmap = {h: i for i, h in enumerate(hdr)}
        for row in ws.iter_rows(values_only=True):
            if row[cmap['label']] in (None, 'label'): continue
            track.add((row[cmap['gene_symbol']], row[cmap['rcd_form_canonical']]))
    # queue
    seen = set(); todo = []
    for x in r:
        g, f = x['gene_symbol'], x['rcd_form_canonical']
        if (g, f) in track or (g, f) in seen: continue
        seen.add((g, f))
        if (x.get('elicit_verbatim_sentence') or '').strip() or (x.get('deep_best_sentence') or '').strip(): continue
        doi = (x.get('consensus_doi') or '').strip()
        if doi: todo.append((g, f, doi))
    # load progress
    prog = {}
    if os.path.exists(PROGRESS):
        prog = json.load(open(PROGRESS, encoding='utf-8'))
    todo = [t for t in todo if t[0] + '|' + t[1] not in prog]
    if limit: todo = todo[:limit]
    print(f'to fetch: {len(todo)}', flush=True)
    sess = requests.Session()
    sess.headers['User-Agent'] = 'RCDome/1.0 (audit)'
    done = 0; ok = 0
    for g, f, doi in todo:
        key = f'{g}|{f}'
        ab = fetch_abstract(doi, sess)
        if ab == '__ERR__':
            prog[key] = {'gene': g, 'form': f, 'doi': doi, 'verbatim': '', 'status': 'FETCH_ERROR'}
        elif ab == '':
            prog[key] = {'gene': g, 'form': f, 'doi': doi, 'verbatim': '', 'status': 'NO_ABSTRACT'}
        else:
            vb = extract_verbatim(ab, g, f)
            prog[key] = {'gene': g, 'form': f, 'doi': doi, 'verbatim': vb, 'status': 'VERBATIM' if vb else 'NO_COOCCUR'}
            if vb: ok += 1
        done += 1
        if done % 50 == 0:
            json.dump(prog, open(PROGRESS, 'w', encoding='utf-8'))
            print(f'  {done}/{len(todo)} done, verbatim ok={ok}', flush=True)
        time.sleep(0.5)
    json.dump(prog, open(PROGRESS, 'w', encoding='utf-8'))
    # write CSV
    with open(OUT, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['gene_symbol', 'rcd_form_canonical', 'consensus_doi', 'verbatim', 'status'])
        for key, v in prog.items():
            w.writerow([v['gene'], v['form'], v['doi'], v['verbatim'], v['status']])
    print(f'DONE batch: {done} fetched, {ok} with verbatim, saved to {OUT}', flush=True)

if __name__ == '__main__':
    lim = int(sys.argv[1]) if len(sys.argv) > 1 else None
    main(lim)

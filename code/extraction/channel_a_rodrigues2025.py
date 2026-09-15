"""
channel_a_rodrigues2025.py — Extract genes from Rodrigues et al. 2025
=======================================================================
Explodes slash-separated RCD form annotations from Dataset S1B_FerrDBv3_merged_v2.xlsx
into one-gene-per-RCD-form-per-row format with source traceability.
Output: gene_rcd_rodrigues2025.csv
Reference: Rodrigues de Souza et al. (2025) Frontiers in Bioinformatics
"""

import openpyxl, csv, os, sys
from collections import defaultdict

INPUT = os.path.join(os.path.dirname(__file__), '..', '..', 'Manuscript', 'Dataset S1B_FerrDBv3_merged_v2.xlsx')
OUTPUT = os.path.join(os.path.dirname(__file__), '..', 'output databases', 'gene_rcd_rodrigues2025.csv')

def extract():
    wb = openpyxl.load_workbook(INPUT, read_only=True, data_only=True)
    ws = wb['Dataset S1B']
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    col_map = {h: i for i, h in enumerate(headers)}

    accepted = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        gene = str(row[col_map['Gene symbol']]).strip() if row[col_map['Gene symbol']] else ''
        if not gene or gene == 'None': continue
        rcd_raw = str(row[col_map['RCD form']]).strip() if row[col_map['RCD form']] else ''
        rcddb = str(row[col_map['RCDdb Matched gene']]).strip()
        ferr_matched = str(row[col_map['FerrDB_v3_Matched']]).strip() if row[col_map.get('FerrDB_v3_Matched', '')] else ''

        if not rcd_raw: continue
        rcd_forms = [f.strip() for f in rcd_raw.split('/') if f.strip()]

        for rcd in rcd_forms:
            sources = ['PubMed_Boolean']
            if ferr_matched == 'Yes': sources.append('FerrDB')
            if rcddb == 'Yes': sources.append('RCDdb')
            channel = 'A' if len(sources) == 1 else 'A,B'

            accepted.append([gene, rcd, ';'.join(sources), channel])

    wb.close()

    with open(OUTPUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['gene_symbol','rcd_form','source_databases','source_channel','evidence_type','curator_verified'])
        for r in accepted:
            w.writerow([r[0], r[1], r[2], r[3], 'pubmed_boolean', 'FALSE'])
    print(f'Rodrigues 2025: {len(accepted)} entries -> {OUTPUT}')

if __name__ == '__main__':
    extract()

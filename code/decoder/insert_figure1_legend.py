# -*- coding: utf-8 -*-
import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript/RCDome_manuscript_governance_V024.docx')

legend = ("Figure 1. RCDome compilation and validation pipeline. RCDome is assembled in a single "
"top-to-bottom workflow. Twenty source databases \u2014 ten RCD-specific curated resources (FerrDB v3, "
"RCDdb, iPCD, LncRCD, iLIR, HADb, HAMdb, ATD, AMT-DB and ncFO), five pathway and molecular databases "
"(KEGG, Reactome, Gene Ontology, MSigDB/GSEA and WikiPathways), NCBI Gene, Rodrigues et al. (2025), "
"XDeathDB, and two Elicit channels (Phase 3 systematic review and Alert) \u2014 are compiled through four "
"complementary channels: Channel A (PubMed Boolean searches) maximizes recall, Channel B (curated "
"databases) provides precision, Channel C (AI-assisted screening) captures recent discoveries, and "
"Channel D (manual curation) serves as the final arbiter. Consolidation normalizes 62 original name "
"strings into 53 canonicalized RCD-related entities, producing 22,473 gene\u2013RCD rows spanning 9,265 unique human "
"genes. The consolidated matrix then enters an independent validation stage, in which three mutually "
"independent analyses run in execution order: (1) the DeepSeek chain-of-thought classifier (Phases 2A/D, "
"2B, 2C, 2D/2E and 2F), which applies two-call validation, deep retrieval-augmented full-text recovery "
"and a full-matrix chain-of-thought classification returning four columns (cot_role, cot_direction, "
"cot_mechanistic, cot_evidence); (2) the Elicit systematic-review extraction (Phase 3), which assigns "
"functional role and direction; and (3) the Consensus evidence-channel classification (Phase 2G), which "
"queries each pair with a neutral \u201cgene cell death\u201d formulation. All three analyses are kept in "
"separate columns so that provenance is never conflated. The final consolidated matrix (22,473 rows "
"\u00d7 71 columns) powers the RCDome web application, whose interface comprises the interactive dome, "
"Explore RCDome, Gene Catalog, RCD entities, Network and RCDchat. Boxes are color-coded for visual "
"distinction; arrows denote the flow of data through the pipeline.")

# find the Figure 3 legend paragraph and insert Figure 1 before it
target = None
for p in doc.paragraphs:
    if p.text.startswith('Figure 3.'):
        target = p
        break

if target is None:
    print('ERROR: Figure 3 paragraph not found')
    sys.exit(1)

new_p = target.insert_paragraph_before()
run = new_p.add_run(legend)
run.bold = True

doc.save('D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript/RCDome_manuscript_governance_V024.docx')
print('Figure 1 legend inserted before Figure 3 (bold, Normal style).')

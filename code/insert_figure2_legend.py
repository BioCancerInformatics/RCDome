# -*- coding: utf-8 -*-
import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript/RCDome_manuscript_governance_V024.docx')

legend2 = ("Figure 2. Benchmark evaluation of the three independent analyses. (A) DeepSeek "
"model-selection benchmark against a fixed 26-pair ground-truth set (18 positive executor/regulator pairs "
"and 8 form-mismatch negative controls): the baseline deepseek-chat model reached 81% combined accuracy "
"with 38% false positives, chain-of-thought deepseek-reasoner reached 73% with 25%, whereas the "
"governance-augmented tool-call loop with a hand-curated reference and the self-provided chain-of-thought "
"loop each reached 100% combined accuracy with zero false positives; the self-provided loop was adopted "
"because it requires no external reference. (B) Elicit post-export governance benchmark on the 784 records "
"extracted across three forms (393 Ferroptosis, 390 Cuproptosis, 1 Triaptosis): allow-list compliance "
"improved from 75.4% to 100%, EXECUTOR precision from 47% to 100% (192 EXECUTOR-to-REGULATOR downgrades), "
"and PMID traceability from 24% to 98%. (C) Consensus query-formulation benchmark against the same 26-pair "
"ground-truth set: the form-biased query reached 92% combined accuracy with 12% false positives, whereas "
"the neutral \u201cgene cell death\u201d query reached 100% combined accuracy with zero false positives. "
"Green bars indicate combined accuracy (A, C) or post-governance values (B); salmon bars indicate false "
"positives (A, C) or pre-governance values (B).")

target = None
for p in doc.paragraphs:
    if p.text.startswith('Figure 3.'):
        target = p
        break

if target is None:
    print('ERROR: Figure 3 paragraph not found')
    sys.exit(1)

new_p = target.insert_paragraph_before()
run = new_p.add_run(legend2)
run.bold = True

doc.save('D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript/RCDome_manuscript_governance_V024.docx')
print('Figure 2 legend inserted before Figure 3 (bold, Normal style).')

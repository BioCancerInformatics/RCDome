# -*- coding: utf-8 -*-
"""Fix supplementary-tables book: correct S2-S7 titles, S12 content, add S13, rebuild INDEX."""
import openpyxl, docx, pandas as pd

BASE = r"D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript"
BOOK = BASE + "/RCDome_manuscript_V023_supplementary_tables.xlsx"

wb = openpyxl.load_workbook(BOOK)

# ---- 1. Fix 6 scrambled titles (title number -> sheet number) ----
fixes = {
    'Table S2': 'Table S7',
    'Table S3': 'Table S6',
    'Table S4': 'Table S2',
    'Table S5': 'Table S3',
    'Table S6': 'Table S4',
    'Table S7': 'Table S5',
}
for sheet, wrong in fixes.items():
    ws = wb[sheet]
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith(wrong):
                cell.value = v.replace(wrong, sheet, 1)
                break
        else:
            continue
        break
print('1. titles fixed')

# ---- 2. Replace S12 content with the 25-forms reference table ----
doc = docx.Document(BASE + "/RCD_25_Forms_Reference_Table_FINAL.docx")
t = doc.tables[0]
ws12 = wb['Table S12']
for row in ws12.iter_rows():
    for cell in row:
        cell.value = None
ws12.cell(row=1, column=1, value='Table S12: RCD 25-Forms Reference Table (20 executioner, 2 trigger, 3 functional)')
for i, row in enumerate(t.rows, start=2):
    for j, cell in enumerate(row.cells, start=1):
        ws12.cell(row=i, column=j, value=cell.text.strip())
print('2. S12 replaced with 25-forms table (rows:', len(t.rows), ')')

# ---- 3. Add S13 = Rodrigues gene-RCD classification ----
df = pd.read_csv(BASE + "/output databases/gene_rcd_rodrigues2025.csv", dtype=str, keep_default_na=False)
if 'Table S13' in wb.sheetnames:
    del wb['Table S13']
ws13 = wb.create_sheet('Table S13')
ws13.cell(row=1, column=1, value='Table S13: Rodrigues 2025 Gene-RCD Classification (rodr_category)')
for j, col in enumerate(df.columns, start=1):
    ws13.cell(row=3, column=j, value=col)
for i, (_, row) in enumerate(df.iterrows(), start=4):
    for j, col in enumerate(df.columns, start=1):
        ws13.cell(row=i, column=j, value=row[col])
print('3. S13 added (rows:', len(df), ')')

# ---- 4. Rebuild INDEX ----
ws_idx = wb['INDEX']
# read existing descriptions (keyed by correct title)
desc = {}
for row in ws_idx.iter_rows(values_only=True):
    vals = [v for v in row if v is not None]
    if vals and str(vals[0]).startswith('Table S'):
        key = str(vals[0]).split(':')[0].strip()  # e.g. "Table S4"
        desc[key] = list(row)

# clear INDEX
for row in ws_idx.iter_rows():
    for cell in row:
        cell.value = None

ws_idx.cell(row=1, column=1, value='INDEX — Supplementary Tables S1–S13')
ws_idx.cell(row=3, column=1, value='Table')
ws_idx.cell(row=3, column=2, value='Title')
ws_idx.cell(row=3, column=3, value='Comprehensive description')
ws_idx.cell(row=3, column=4, value='Dimensions')

order = ['Table S1','Table S2','Table S3','Table S4','Table S5','Table S6','Table S7',
         'Table S8','Table S9','Table S10','Table S11']
r = 4
for key in order:
    vals = desc.get(key, [key, '', '', ''])
    ws_idx.cell(row=r, column=1, value=vals[0])
    ws_idx.cell(row=r, column=2, value=vals[1] if len(vals) > 1 else '')
    ws_idx.cell(row=r, column=3, value=vals[2] if len(vals) > 2 else '')
    ws_idx.cell(row=r, column=4, value=vals[3] if len(vals) > 3 else '')
    r += 1
# S12 + S13
ws_idx.cell(row=r, column=1, value='Table S12')
ws_idx.cell(row=r, column=2, value='RCD 25-Forms Reference Table')
ws_idx.cell(row=r, column=3, value='The 25 recognized RCD forms (20 executioner, 2 trigger, 3 functional) with category, core mechanism, key molecules, inflammatory status, cancer relevance, and PMID references.')
ws_idx.cell(row=r, column=4, value='26 rows x 9 cols'); r += 1
ws_idx.cell(row=r, column=1, value='Table S13')
ws_idx.cell(row=r, column=2, value='Rodrigues 2025 Gene-RCD Classification')
ws_idx.cell(row=r, column=3, value='External Rodrigues 2025 classification of gene-RCD associations (rodr_category), preserved untouched.')
ws_idx.cell(row=r, column=4, value=f'{len(df)} rows x {len(df.columns)} cols')

wb.save(BOOK)
print('4. INDEX rebuilt; saved.')

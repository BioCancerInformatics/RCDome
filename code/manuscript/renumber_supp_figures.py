# -*- coding: utf-8 -*-
"""Renumber supplementary figures S1<->S3 (infographic becomes S1) in the V026 docx files.
Order of operations (reverse, to avoid collisions): S2->S3, then S1->S2, then insert new S1."""
import sys, shutil, datetime
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

BASE = r"D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript/manuscript_V027"
GOV = f"{BASE}/RCDome_manuscript_governance_V026.docx"
SUP = f"{BASE}/RCDome_manuscript_V027_supplementary_information.docx"
stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')

# backups (no overwrite)
shutil.copy(GOV, f"{GOV.replace('.docx','')}_backup_S3renumber_{stamp}.docx")
shutil.copy(SUP, f"{SUP.replace('.docx','')}_backup_S3renumber_{stamp}.docx")

NEW_S1_LEGEND = """Supplementary Figure S1. Evaluation architecture of RCDome. RCDome decomposes its curation into three objects that answer distinct scientific questions and are therefore assessed independently and never conflated. (1) For the RCD entity, the question is: “What kind of cell-death entity is this?” The answer is provided by modality criteria (MC1–MC5: machinery, trigger, morphology, nomenclature, distinguishability), an evidence checklist (Q1–Q13), biological gates (BG0–BG6, with G1 machinery and G2 necessity obligatory), and a final ontological level (L1–L5); complementary descriptors record evidence maturity, species scope, relationship type, and parent entity. (2) For the gene–RCD association, the question is: “Is gene X mechanistically linked to this entity?” The answer is provided by evidence tier (ET1–ET4) and a distinct confidence tier (CT1–CT4 or confidence_score), with mechanistic role and direction (EXECUTOR, REGULATOR, MARKER; positive/negative/mixed). (3) For pipeline quality, the question is: “Can the automated extraction be trusted?” The answer is provided by Elicit governance gates (EG1–EG5), input validation (allow-list and vocabulary compliance), EXECUTOR precision, and traceability. This three-way separation ensures that an ontological judgment is never mistaken for gene-level evidence, and that an extraction error is never mistaken for biological signal."""


def para_replace(par, old, new):
    full = ''.join(r.text for r in par.runs)
    if old in full:
        newfull = full.replace(old, new)
        if par.runs:
            par.runs[0].text = newfull
            for r in par.runs[1:]:
                r.text = ''
        return 1
    return 0

# ---------- 1. GOVERNANCE manuscript ----------
doc = Document(GOV)
n = 0
for par in doc.paragraphs:
    n += para_replace(par, "(Supplementary Figure S2)", "(Supplementary Figure S3)")   # source -> S3
    n += para_replace(par, "(Supplementary Figure S1)", "(Supplementary Figure S2)")   # database -> S2
    n += para_replace(par,
        "The adjudication framework and its governance are documented in Supplementary Note S1.",
        "The adjudication framework and its governance are documented in Supplementary Note S1; the complete "
        "evaluation architecture \u2014 RCD entity classification, gene\u2013RCD association evidence, and pipeline "
        "quality gates \u2014 is summarized in Supplementary Figure S1.")
doc.save(GOV)
print(f"[governance] replacements applied: {n}")

# ---------- 2. SUPPLEMENTARY INFORMATION ----------
doc2 = Document(SUP)
m = 0
# renumber the two existing legends (reverse order)
for par in doc2.paragraphs:
    m += para_replace(par, "Supplementary Figure S2. Mechanistic associations", "Supplementary Figure S3. Mechanistic associations")
    m += para_replace(par, "Figure S1. Database contribution", "Supplementary Figure S2. Database contribution")
# insert the new S1 legend before the (now) S2 legend paragraph
for i, par in enumerate(doc2.paragraphs):
    if par.text.startswith("Supplementary Figure S2. Database contribution"):
        # copy style from the existing legend paragraph
        newpar = par.insert_paragraph_before(NEW_S1_LEGEND)
        newpar.style = par.style
        m += 1
        break
doc2.save(SUP)
print(f"[supplementary] replacements applied: {m}")
print("backups:", stamp)

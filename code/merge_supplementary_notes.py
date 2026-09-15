# -*- coding: utf-8 -*-
import docx, copy, sys
sys.stdout.reconfigure(encoding='utf-8')
from docx.oxml.ns import qn

BASE = 'D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript/'
SUPP = BASE + 'RCDome_manuscript_V024_supplementary_information.docx'

def copy_note_into(target, note_path):
    src = docx.Document(note_path)
    body = target.element.body
    sectPr = body.find(qn('w:sectPr'))
    for child in src.element.body.iterchildren():
        if child.tag in (qn('w:p'), qn('w:tbl')):
            sectPr.addprevious(copy.deepcopy(child))

target = docx.Document()
target.add_heading('RCDome \u2014 Supplementary Information', level=0)

# ---- Supplementary Notes (first) ----
target.add_heading('Supplementary Notes', level=1)
copy_note_into(target, BASE + 'RCDome_AI_usage_supplementary_note_S1.docx')
copy_note_into(target, BASE + 'RCDome_CoT_benchmark_supplementary_note_S2.docx')
copy_note_into(target, BASE + 'RCDome_Consensus_benchmark_supplementary_note_S3.docx')

# ---- Supplementary Figure Legends (second) ----
target.add_heading('Supplementary Figure Legends', level=1)
figS1 = ("Supplementary Figure S1. Database contribution and multi-database corroboration. (A) Number of "
"gene\u2013RCD associations contributed by each of the 20 source databases (log2-scaled axis; exact counts printed "
"at the end of each bar). (B) Number of source databases per association: 10,009 associations rest on a single "
"source, 6,313 are corroborated by two sources, and the count declines steeply with additional sources; the long "
"tail up to 14 sources marks the best-corroborated associations (e.g., the canonical executors), which carry the "
"strongest multi-database agreement.")
figS2 = ("Supplementary Figure S2. Fraction of associations with mechanistic classification, per source. Each bar "
"shows the percentage of a source\u2019s gene\u2013RCD associations classified as mechanistically involved by the DeepSeek "
"chain-of-thought run (cot_mechanistic = YES), with the numerator and denominator printed at the end of each bar. "
"Classification yield ranges from 21.7% (XDeathDB, 378 of 1,739) to 90.0% (ncFO, 45 of 50), revealing which "
"sources contribute mostly mechanistically-validated associations.")
p1 = target.add_paragraph(); p1.add_run(figS1).bold = True
p2 = target.add_paragraph(); p2.add_run(figS2).bold = True

target.save(SUPP)
print('merged REAL S1/S2/S3 note files + figure legends')

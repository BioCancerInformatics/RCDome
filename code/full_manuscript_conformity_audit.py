import pandas as pd, re, os, sys, json, hashlib, subprocess, zipfile
import docx, openpyxl
sys.stdout.reconfigure(encoding='utf-8')
B = r"D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript"
MIR = r"D:/RCDome_github"
P = F = 0
def ck(n, ok, d=''):
    global P, F
    print(f'  [{"PASS" if ok else "FAIL"}] {n} {d}')
    P += bool(ok); F += (not bool(ok))

EXPECT = {'L1':16,'L2':5,'L3':2,'L4':12,'L5':18}

print('='*74)
print('FULL MANUSCRIPT CONFORMITY AUDIT — every single artifact')
print('='*74)

# ── 1. MATRIX ──
print('\n[1] MATRIX V026')
m = pd.read_csv(B+'/output databases/Phase 2C output/unified_gene_rcd_matrix_V026.csv', low_memory=False)
ck('22,473 x 71', m.shape==(22473,71), str(m.shape))
ck('levels 16/5/2/12/18', m.drop_duplicates('rcd_form_canonical').rcdome_level.value_counts().to_dict()==EXPECT)
ck('decoder_* 4 cols', all(c in m.columns for c in ['decoder_operational_definition','decoder_key_molecules','decoder_category','decoder_hard_gate_status']))
ck('53 forms', m.rcd_form_canonical.nunique()==53)
ck('9,265 genes', m.gene_symbol.nunique()==9265)
ck('21,568 pairs', len(m.drop_duplicates(['gene_symbol','rcd_form_canonical']))==21568)

# ── 2. DECODER + D1 ──
print('\n[2] DECODER + D1')
d = pd.read_csv(B+'/RCDome Decoder/RCDome_Decoder_53.csv')
ck('decoder 53x39', d.shape==(53,39), str(d.shape))
ck('defs 53/53', d.operational_definition.notna().all() and (d.operational_definition.astype(str).str.strip()!='').all())
ck('hgs 16/37', (d.hard_gate_status=='L1-ELIGIBLE').sum()==16 and (d.hard_gate_status=='NOT-L1').sum()==37)
d1 = pd.read_csv(B+'/RCDome Decoder/D1_adjudication_table.csv')
ck('D1 53x26', d1.shape==(53,26), str(d1.shape))

# ── 3. MANUSCRIPT ──
print('\n[3] MANUSCRIPT')
doc = docx.Document(B+'/manuscript_V026/RCDome_manuscript_governance.docx')
paras = [p.text for p in doc.paragraphs]
txt = ' '.join(paras)
ck('version-free', not re.search(r'V0\d\d', txt))
ck('71 columns (no 67)', '71 columns' in txt and '67 columns' not in txt and '67-column' not in txt)
ck('Database S1 (no Dataset)', 'Database S1' in txt and 'Dataset S1' not in txt)
ck('16 modalities (no 38)', '16 RCD modalities' in txt and '38 modalities' not in txt and '38 RCD modalities' not in txt)
ck('no old labels', 'functional/conceptual' not in txt and 'generic/composite' not in txt)
ck('hard gates', 'hard gates' in txt and 'G0' in txt)
ck('D2 validations', 'Two distinct validations' in txt)
ck('F2a/F2b/F3abc', all(s in txt for s in ['not be read as an independent validation','database indexing, not ontological inheritance','RCD-related entities','dynamic, not binary','curator-governed']))
ck('Table S12 callout', 'documented in Table S12' in txt)
ck('stats 22,473/21,568/9,265', all(s in txt for s in ['22,473','21,568','9,265']))

# ── 4. FIGURE LEGENDS (inside manuscript) ──
print('\n[4] FIGURE LEGENDS')
for i in range(355, 370):
    t = paras[i].strip()
    if t.startswith('Figure 1.') or t.startswith('Figure 2.') or t.startswith('Figure 3.') or t.startswith('Figure 4.'):
        n = t.split('.')[0]
        ok = True; why=''
        if 'Figure 1' in n:
            ok = '71 columns' in t and '22,473' in t
            why = 'Fig1 71/22,473'
        elif 'Figure 3' in n:
            ok = '21,568' in t and '9,489 regulators' in t
            why = 'Fig3 21,568/9,489'
        elif 'Figure 4' in n:
            ok = 'L4 Functional/contextual' in t and 'L5 Generic/umbrella/legacy/taxon/emerging' in t
            why = 'Fig4 labels'
        ck(f'legend {n}: {why}', ok, '')

# ── 5. FIGURE FILES ──
print('\n[5] FIGURE FILES')
for f in ['Figure_1_pipeline','Figure_2_benchmarks','Figure_3','Figure_4_entity_coverage','Figure_S2_database_contribution','Figure_S3_source_classification']:
    tiff = os.path.exists(B+f'/Figures/{f}.tiff')
    ck(f'{f}.tiff', tiff)

# ── 6. SUPPLEMENTARY WORKBOOK ──
print('\n[6] SUPPLEMENTARY WORKBOOK')
wb = openpyxl.load_workbook(B+'/manuscript_V026/RCDome_manuscript_supplementary_tables.xlsx', read_only=True)
ck('sheets S1-S15', all(f'Table S{i}' in wb.sheetnames for i in range(1,16)))
ws = wb['Table S12']; hdr=[c.value for c in next(ws.iter_rows(max_row=1))]
ck('S12 26 cols', len(hdr)==26 and 'adjudicated_level' in hdr)
md_ws = wb['Matrix Descriptor']
md_rows = sum(1 for row in md_ws.iter_rows() if row[1].value and row[1].value!='Variable')
ck('Matrix Descriptor 71 variables', md_rows==71, f'{md_rows}')
idx = ' '.join(str(c.value) for row in wb['INDEX'].iter_rows() for c in row if c.value)
ck('INDEX S1-S15 + Database S1', 'S1–S15' in idx and 'Database S1' in idx)

# ── 7. SUPPLEMENTARY INFORMATION ──
print('\n[7] SUPPLEMENTARY INFORMATION')
si = docx.Document(B+'/manuscript_V026/RCDome_manuscript_supplementary_information.docx')
sit = ' '.join(p.text for p in si.paragraphs)
ck('SI version-free', not re.search(r'V0\d\d', sit))
ck('SI has Notes S1-S3', all(n in sit for n in ['Note S1','Note S2','Note S3']))

# ── 8. APP ──
print('\n[8] APP')
APP = B+'/Shiny app/RCDome_Atlas/APP'
import pyreadr
r = pyreadr.read_r(APP+'/data/unified_gene_rcd_matrix_V026.rds')
kk = list(r.keys())[0]
ck('app RDS 71', r[kk].shape==(22473,71))
for f in ['www/landing_page.html','www/dome_prototype.html']:
    t = open(APP+'/'+f, encoding='utf-8').read()
    lc = re.search(r'LEVEL_COUNTS = \{([^}]*)\}', t)
    counts = dict(re.findall(r'"(L\d)"\s*:\s*(\d+)', lc.group(1))) if lc else {}
    ck(f'{f.split("/")[-1]} counts', all(counts.get(k)==str(v) for k,v in EXPECT.items()), str(counts))
t = open(APP+'/www/detail_page_wireframe.html', encoding='utf-8').read()
i = t.find('"entities"'); start=t.find('{',i); depth=0; end=None
for k in range(start,len(t)):
    if t[k]=='{': depth+=1
    elif t[k]=='}':
        depth-=1
        if depth==0: end=k+1; break
ents = json.loads(t[start:end])
dm = [nm for nm in d.rcd_form if nm not in ents or str(ents[nm].get('definition','')).strip()!=str(d.set_index('rcd_form').loc[nm,'operational_definition']).strip()]
ck('detail defs verbatim 53/53', len(dm)==0, f'{len(dm)}')
ex = open(APP+'/www/explore.html', encoding='utf-8').read()
ck('explore stats', all(s in ex for s in ['21,568','469','9,489','Human ontological curation']))
ck('explore no stale 67-column', '67 columns' not in ex and '67-column' not in ex)

# ── 9. ZIP ──
print('\n[9] ZIP PACKAGE')
zpath = B+'/Shiny app/RCDome_Atlas/APP/RCDome_V026_update.zip'
z = zipfile.ZipFile(zpath); names = z.namelist()
ck('zip 36 files + key data', len(names)==36 and 'data/unified_gene_rcd_matrix_V026.rds' in names and 'data/RCDome_Decoder_53.csv' in names and 'CHANGELOG_V026.md' in names)

# ── 10. GITHUB ──
print('\n[10] GITHUB')
def sha256(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for c in iter(lambda: f.read(65536), b''): h.update(c)
    return h.hexdigest()
mism=0
for l in open(MIR+'/mirror_manifest.txt', encoding='utf-8').read().splitlines():
    if not l.strip(): continue
    parts=l.split('  ')
    if len(parts)!=2: continue
    h,rel=parts; p=os.path.join(MIR, rel.replace('/',os.sep))
    if os.path.exists(p) and sha256(p)!=h: mism+=1
ck('manifest 0 mismatch', mism==0, f'{mism}')
st = subprocess.run(['git','-C',MIR,'status','--porcelain'], capture_output=True, text=True).stdout.strip()
ck('git clean', st=='', st[:80] if st else '')
stale = [f for r,dd,fs in os.walk(MIR) if '.git' not in r for f in fs if ('V025' in f or 'V024' in f or 'V023' in f) and 'merge_cot' not in f]
ck('no stale V02x (except hist code)', len(stale)==0, str(stale))

print(f'\n===== FULL MANUSCRIPT CONFORMITY: {P} PASS, {F} FAIL =====')

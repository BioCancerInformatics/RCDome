# -*- coding: utf-8 -*-
import os, re, sys, io
import docx, pandas as pd
sys.stdout.reconfigure(encoding='utf-8')

BASE = "D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript"
os.chdir(BASE)

out = io.StringIO()
def log(*a): print(*a, file=out)

CANON = re.compile(r"canon", re.IGNORECASE)
# safe (keep) categories
BIO = ["canonical executors", "canonical gene set", "canonical apoptosis", "canonical autophagy",
       "canonical inflammasome", "canonical ferroptosis", "canonical notch", "canonical bmp",
       "canonical nlrp", "canonical signaling", "canonical wnt", "canonical mitochondrial",
       "canonical mapk", "canonical macroautophagy", "canonical cellular stress",
       "canonical atg", "canonical pathway", "canonical mechanism", "canonical vitamin",
       "canonical gsdmd", "canonical caspas", "canonical pd", "canonical hsp", "canonical rack",
       "canonical ubc", "canonical s6k", "canonical lc3", "canonical mek", "non-canonical",
       "noncanonical", "canonical reference framework"]
SCHEMA = ["rcd_form_canonical", "canonical_pathways", "canonical_set", "is_canonical",
          "canonical_symbol", "canonical_transcript", "canonical universe", "canonical gene ",
          "canonical name accuracy", "canonical count", "canonical audit", "audit_canonical",
          "ferrdb_to_canonical", "canonical_rcd_forms", "canonical_databases", "canonical rcd pathway",
          "canonical entity)", "canonical name or its", "canonical values", "canonical concept name",
          "canonical label", "group_g_canonical", "canonical", "canonical_umbrella"]

def classify(frag):
    f = frag.lower()
    if "canonicalized" in f or "canonicalization" in f or "normalized" in f:
        return "FIXED(term applied)"
    if any(b in f for b in BIO):
        return "BIOLOGICAL(keep)"
    if any(s in f for s in SCHEMA):
        return "SCHEMA/CODE(keep)"
    return "REVIEW"

log("=" * 78)
log("OCEAN-DEEP 'canon' AUDIT — final residual classification")
log("=" * 78)

# ---- 1. Manuscript + supplementary docx ----
log("\n## 1. MANUSCRIPT + SUPPLEMENTARY DOCX")
for f in ["manuscript_V026/RCDome_manuscript_governance_V026.docx",
          "manuscript_V026/RCDome_manuscript_V026_supplementary_information.docx"]:
    d = docx.Document(f)
    txt = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                txt += "\n" + c.text
    pat = re.compile(r"[^\n]{0,50}canon[^\n]{0,50}", re.IGNORECASE)
    seen = set()
    for m in pat.finditer(txt):
        frag = " ".join(m.group(0).split())
        k = frag.lower()
        if k not in seen:
            seen.add(k)
            log(f"  [{classify(frag)}] …{frag}…")
    log(f"  (file: {f.split('/')[-1]} — unique: {len(seen)})")

# ---- 2. Supplementary tables xlsx ----
log("\n## 2. SUPPLEMENTARY TABLES XLSX")
x = pd.ExcelFile("manuscript_V026/RCDome_manuscript_V026_supplementary_tables.xlsx")
for sh in x.sheet_names:
    df = x.parse(sh, dtype=str)
    seen = set()
    for col in df.columns:
        for val in df[col]:
            if isinstance(val, str) and "canon" in val.lower():
                frag = " ".join(val.split())
                k = frag.lower()
                if k not in seen:
                    seen.add(k)
                    log(f"  [{sh} | {classify(frag)}] …{frag[:110]}…")
    if seen:
        log(f"  (sheet {sh} — unique: {len(seen)})")

# ---- 3. Matrix V-26 ----
log("\n## 3. MATRIX V-26 (unified_gene_rcd_matrix_V026.csv)")
m = pd.read_csv("output databases/Phase 2C output/unified_gene_rcd_matrix_V026.csv",
                dtype=str, low_memory=False, nrows=200000)
cols = [c for c in m.columns if "canon" in c.lower()]
log("  'canon' columns:", cols)
for c in m.columns:
    if m[c].astype(str).str.contains("canon", case=False, na=False).any():
        vals = sorted(m[c].astype(str)[m[c].astype(str).str.contains("canon", case=False, na=False)].unique())
        log(f"  col '{c}' distinct canon values ({len(vals)}): {vals[:12]}")

# ---- 4. Decoder CSV/xlsx ----
log("\n## 4. DECODER CSV/XLSX")
for f in ["RCDome Decoder/RCDome_Decoder_53.csv"]:
    d = pd.read_csv(f, dtype=str)
    seen = set()
    for col in d.columns:
        for val in d[col]:
            if isinstance(val, str) and "canon" in val.lower():
                frag = " ".join(val.split())
                k = frag.lower()
                if k not in seen:
                    seen.add(k)
                    log(f"  [{col} | {classify(frag)}] …{frag[:110]}…")

# ---- 5. App modules + www + board + code (current) ----
log("\n## 5. APP + CODE (current user-facing + scripts)")
KEY = ["Shiny app/RCDome_Atlas/APP/modules/", "Shiny app/RCDome_Atlas/APP/R/",
       "Shiny app/RCDome_Atlas/APP/global.R", "Shiny app/RCDome_Atlas/APP/server.R",
       "Shiny app/RCDome_Atlas/APP/www/", "Shiny app/RCDome_Atlas/APP/board/",
       "RCDome Decoder/", "code/", "docs/"]
files = []
for root, dirs, fs in os.walk("."):
    dirs[:] = [d for d in dirs if d not in ("__pycache__", ".git", "_archive")]
    for f in fs:
        if not f.endswith((".md", ".R", ".html", ".py", ".txt", ".json")):
            continue
        p = os.path.join(root, f).replace("\\", "/").lstrip("./")
        if any(p == k.rstrip("/") or p.startswith(k) for k in KEY):
            files.append(p)
files = sorted(set(files))
seen = set()
for p in files:
    try:
        txt = io.open(p, encoding="utf-8", errors="replace").read()
    except Exception:
        continue
    for m in re.finditer(r"[^\n]{0,45}canon[^\n]{0,45}", txt, re.IGNORECASE):
        frag = " ".join(m.group(0).split())
        k = (p, frag.lower())
        if k not in seen:
            seen.add(k)
            c = classify(frag)
            if c == "REVIEW":
                log(f"  [REVIEW] {p}: …{frag}…")
log("  (files scanned:", len(files), "| REVIEW-flagged:", sum(1 for k in seen if classify(k[1])=='REVIEW'), ")")

io.open("D:/pa0_ocean.txt", "w", encoding="utf-8").write(out.getvalue())
print(out.getvalue())

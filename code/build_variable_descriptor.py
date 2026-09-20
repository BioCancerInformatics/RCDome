# LEGACY (V025): superseded. Do NOT re-run against V026 — it targets the V025 workbook and its S12-S14 descriptor blocks are off-by-one.
# The V026 Variable Descriptor sheet is maintained directly in RCDome_manuscript_V026_supplementary_tables.xlsx.
"""
Build a consolidated "Variable Descriptor" sheet covering ALL tables S1-S14,
following the logic of Rodrigues_et_al_Supplementary_Tables (Variable Name | Short Descriptor).
Three columns: Table | Variable Name | Short Descriptor.
"""
import openpyxl, sys

BASE = r"D:/Pré-artigo 5-optosis model/Cellular_Spatial_RCD_manuscript/RCDome_manuscript"
WB = BASE + "/RCDome_manuscript_V025_supplementary_tables.xlsx"

DESC = [
 # Table S1
 ("S1","Feature","Resource feature being compared"),
 ("S1","FerrDB v3","FerrDB v3 value for the feature"),
 ("S1","RCDdb","RCDdb value for the feature"),
 ("S1","RCDome (this work)","RCDome value for the feature"),
 # Table S2
 ("S2","#","Rank order"),
 ("S2","RCD Form","Canonicalized RCD-related entity name"),
 ("S2","Category","Executioner / Trigger / Functional"),
 ("S2","Core Mechanism","Short core mechanism"),
 ("S2","Key Molecules","Key molecular machinery"),
 ("S2","Inflammatory?","Inflammatory status (Yes / No / context-dependent)"),
 ("S2","Cancer Relevance","Cancer relevance (Rodrigues 2025)"),
 ("S2","Primary References (PMID)","Primary evidence PMIDs"),
 ("S2","Secondary References (PMID)","Secondary evidence PMIDs"),
 # Table S3
 ("S3","Tier","Database tier (1 / 2 / 3)"),
 ("S3","#","Database rank"),
 ("S3","Database","Database name"),
 ("S3","URL","Database URL"),
 ("S3","Access Method","API / download"),
 ("S3","RCD Forms","RCD forms queried"),
 ("S3","Genes Extracted","Number of genes extracted"),
 # Table S4
 ("S4","#","Rank order"),
 ("S4","RCD Form","Canonicalized RCD-related entity name"),
 ("S4","Category","Category"),
 ("S4","Evidence Grade","Evidence grade"),
 ("S4","Gene Count","Gene count"),
 ("S4","Discovery","Discovery year"),
 ("S4","Key References (PMID)","Key evidence PMIDs"),
 # Table S5
 ("S5","gene_symbol","HGNC gene symbol"),
 ("S5","entrez_id","NCBI Entrez Gene ID"),
 ("S5","rcd_form","RCD form association"),
 ("S5","source_databases","Source databases"),
 ("S5","source_channel","Source channel"),
 ("S5","evidence_type","Evidence type"),
 ("S5","pmid_list","Supporting PMIDs"),
 ("S5","relationship","Gene-RCD relationship"),
 ("S5","gene_type","Gene type"),
 ("S5","evidence_strength","Evidence strength"),
 ("S5","curator_verified","Curator verification"),
 ("S5","rcd_count_total","Total RCD count"),
 ("S5","chromosome","Chromosome"),
 ("S5","map_location","Map location"),
 ("S5","cosmic_driver","COSMIC driver"),
 ("S5","cosmic_tier","COSMIC tier"),
 ("S5","present_in_signature_db","Present in signature database"),
 ("S5","description","Gene description"),
 # Table S6
 ("S6","Original Name","Original name string"),
 ("S6","Normalized Name","Normalized name"),
 ("S6","Rationale","Renaming rationale"),
 # Table S7
 ("S7","Column","Column name"),
 ("S7","Description","Column description"),
 # Table S8
 ("S8","Gate","Benchmark gate (Q1-Q5)"),
 ("S8","Metric","Metric"),
 ("S8","Threshold","Threshold"),
 ("S8","Rationale","Rationale"),
 # Table S9
 ("S9","#","Rank"),
 ("S9","Form","RCD form name"),
 ("S9","Grade","Grade"),
 ("S9","Year","Discovery year"),
 ("S9","Mechanism","Mechanism"),
 ("S9","Flag","Flag"),
 # Table S10
 ("S10","Script","Script name"),
 ("S10","Tier","Tier"),
 ("S10","Source","Source"),
 # Table S11 (decoder 31 vars)
 ("S11","rcd_form","Canonicalized RCD-related entity name (53 entities)"),
 ("S11","category","Entity category (Executioner / Trigger / Functional / Subtype / Synonym / Umbrella / Contextual / Taxon / Cell-type / Differentiation / Legacy / Emerging / Integrated / Composite / Morphology)"),
 ("S11","year","Year of discovery (of the founding reference; 'composite' for non-single entities)"),
 ("S11","operational_definition","Operational definition prose (Galluzzi C1-C5 logic, classification-led, citation-free)"),
 ("S11","operational_definition_complement","Galluzzi-style complement for verbatim entities (never replaces the verbatim)"),
 ("S11","operational_definition_source","Source of the definition (NCCD_2018_VERBATIM / RODRIGUES_S1A / FOUNDING_DRAFT / PRE_PUBMED / COMPOSITE)"),
 ("S11","reference_pmid","PMID anchoring the definition (evidence)"),
 ("S11","reference_citation","Citation (Author et al. Year)"),
 ("S11","primary_refs_pmids","Primary evidence PMIDs (semicolon-separated)"),
 ("S11","secondary_refs_pmids","Secondary evidence PMIDs (semicolon-separated; empty if none)"),
 ("S11","core_mechanism","Short core mechanism (from the 25-Forms Reference Table)"),
 ("S11","key_molecules","Key molecular machinery (comma-separated)"),
 ("S11","inflammatory","Inflammatory status (Yes / No / not specified / member-dependent)"),
 ("S11","overlaps","Description of overlaps with other RCD forms"),
 ("S11","C1","Extent of Research criterion (1-10, weight 25%)"),
 ("S11","C2","Direct Involvement criterion (1-10, weight 25%)"),
 ("S11","C3","Impact on cell/tissue homeostasis criterion (1-10, weight 20%)"),
 ("S11","C4","Emerging Potential criterion (1-10, weight 15%)"),
 ("S11","C5","Biological Significance criterion (1-10, weight 10%)"),
 ("S11","C6","Pathophysiological Relevance criterion (1-10, weight 5%)"),
 ("S11","S","Overall Score = 0.25*C1 + 0.25*C2 + 0.20*C3 + 0.15*C4 + 0.10*C5 + 0.05*C6 (deterministic)"),
 ("S11","founding_pmid","Founding PMID (the discovery article)"),
 ("S11","definition_status","Definition status (DEFINED_RODRIGUES / DRAFT_SOLID / DRAFT_CANDIDATE / COMPOSITE / PRE_PUBMED)"),
 ("S11","galluzzi_nccd_definition","Verbatim NCCD 2018 Box 1 definition (byte-identical; for invariant entities)"),
 ("S11","galluzzi_nccd_status","NCCD status (NCCD_2018_INVARIANT / NOT_IN_NCCD_2018)"),
 ("S11","source_databases","Source databases contributing the entity's genes (semicolon-separated)"),
 ("S11","source_channels","Source channels (A = programmatic / B = database / E = AI-screening)"),
 ("S11","provenance_note","Provenance note (matrix ontology annotation)"),
 ("S11","od_author","Authoring source of the definition (Galluzzi et al. 2018 / Rodrigues et al. 2025 / Rodrigues et al. RCDome)"),
 ("S11","review_status","Review status (APPROVED = human-approved draft / LOCKED-VERBATIM = published verbatim)"),
 ("S11","validation_verdict","Validation verdict (VALIDATED / MINOR_ISSUE / VERBATIM-INVARIANT / CRITERION-PENDING / UMBRELLA-SCOPE / ABSTRACT-LIMITED / LOCKED-VERBATIM)"),
 # Table S12
 ("S12","Form","RCD form"),
 ("S12","Allow-list genes","Allow-list genes"),
 ("S12","Papers screened","Papers screened"),
 ("S12","Raw records","Raw records"),
 ("S12","Post-governance records","Post-governance records"),
 ("S12","Governance actions","Governance actions"),
 # Table S13
 ("S13","Gate","Benchmark gate"),
 ("S13","Metric","Metric"),
 ("S13","Score","Score"),
 ("S13","Threshold","Threshold"),
 ("S13","Result","Pass / Fail result"),
 # Table S14
 ("S14","maxResults","Maximum search results"),
 ("S14","Papers","Papers"),
 ("S14","Governed records","Governed records"),
 ("S14","Unique genes","Unique genes"),
]

wb = openpyxl.load_workbook(WB)
if "Variable Descriptor" in wb.sheetnames:
    del wb["Variable Descriptor"]
if "Decoder Variable Descriptor" in wb.sheetnames:
    del wb["Decoder Variable Descriptor"]
ws = wb.create_sheet("Variable Descriptor")
ws.cell(row=1, column=1, value="Table")
ws.cell(row=1, column=2, value="Variable Name")
ws.cell(row=1, column=3, value="Short Descriptor")
for i, (t, name, desc) in enumerate(DESC, 2):
    ws.cell(row=i, column=1, value=t)
    ws.cell(row=i, column=2, value=name)
    ws.cell(row=i, column=3, value=desc)
# position: INDEX, Variable Descriptor, Table S*, Matrix Descriptor
idx = [w for w in wb.worksheets if w.title == "INDEX"]
vd = [w for w in wb.worksheets if w.title == "Variable Descriptor"]
ts = sorted([w for w in wb.worksheets if w.title.startswith("Table S")], key=lambda w: int(w.title.split("S")[1]))
md = [w for w in wb.worksheets if w.title == "Matrix Descriptor"]
wb._sheets = idx + vd + ts + md
wb.save(WB)
print("SAVED: Variable Descriptor sheet with", len(DESC), "variable entries across S1-S14")
